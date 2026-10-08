"""Recalculate the workbook in LibreOffice and compare ~150 cells against the independent Python engine."""
import subprocess, sys, os, shutil, math
from openpyxl import load_workbook
import engine as E
import build_excel as BX

XL = sys.argv[1] if len(sys.argv) > 1 else "../excel/CDSCO_45kW_Ballast_MMS_Design_Calc.xlsx"
tmp = "/tmp/claude-0/-home-user-Cluade/cb1ede1f-45e1-5aee-bb0b-19ef0d08b619/scratchpad/recalc"
shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, XL], check=True, capture_output=True, timeout=180)
rc = os.path.join(tmp, os.path.basename(XL))
wb = load_workbook(rc, data_only=True)
bk = BX.build(os.path.join(tmp, "_probe.xlsx"))                 # for name -> address map and row registry

def val(name):
    sh, ref = bk.names[name].split("!"); return wb[sh][ref.replace("$", "")].value

# ---- 1. error scan -------------------------------------------------------------------------------------
errs = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (c.value in ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!") or c.value.startswith("Err:")):
                errs.append((ws.title, c.coordinate, c.value))
print("formula errors:", errs[:10] if errs else "NONE", "| sheets:", len(wb.worksheets))

# ---- 2. comparisons ------------------------------------------------------------------------------------
o = E.design(); g, w, L, mc, bal = o["g"], o["w"], o["L"], o["mc"], o["bal"]; p = o["p"]
cmp = []
def add(lbl, xl, py): cmp.append((lbl, xl, py))
for k in ("AC", "AB", "BC", "tilt_deg", "closure", "J1", "J2", "J3", "J4", "s1", "s2", "S1x", "S1y", "S2x", "S2y", "xQ", "yQ", "Ieff", "xpr", "xpl", "xb1", "xb2", "yA", "yB"):
    pk = {"S1x": g["S1"][0], "S1y": g["S1"][1], "S2x": g["S2"][0], "S2y": g["S2"][1], "xQ": bal["xQ"], "yQ": bal["yQ"], "xpr": bal["xp_r"], "xpl": bal["xp_l"], "xb1": bal["xb1"], "xb2": bal["xb2"]}.get(k, g.get(k))
    add("g_" + k, val("g_" + k), pk)
for k in ("z", "k2", "Vz", "pz", "pd", "cp_dn", "cp_up", "Nup", "Ndn"):
    pk = {"Nup": w["N_up"], "Ndn": w["N_dn"]}.get(k, w.get(k)); add("w_" + k, val("w_" + k), pk)
add("l_Ah", val("l_Ah"), L["Ah"]); add("l_ws", val("l_ws"), L["ws"]); add("l_Wm", val("l_Wm"), L["Wm"])
add("sec_Zg", val("sec_Zg"), mc["Zg"]); add("sec_Zn", val("sec_Zn"), mc["Zn"]); add("sec_Ieff", val("sec_Ieff"), g["Ieff"])
# Analysis sheet vs closed form, every LC
an = wb["Analysis"]; ar = bk.reg["Analysis"]
for j, lc in enumerate(BX.LCS):
    res = o["res"][lc]
    for k in ("VB", "Cx", "Cy", "NAB", "qnm", "qnr", "N_0", "N_a", "N_1m", "N_1p", "N_2m", "N_2p", "N_b", "N_L", "V_0", "V_a", "V_1m", "V_1p", "V_2m", "V_2p", "V_b", "V_L",
              "Minc1", "Minc2", "MincM", "M1", "M2", "M3", "M4", "Mm1", "Mm2", "Mm3", "R1", "R2", "R3", "R4", "RH1", "dA", "dC", "V1p", "V3m", "NbA", "NbC", "Fnb", "Ftb"):
        add(f"Analysis {lc} {k}", an.cell(ar[k], 4 + j).value, res[{'NbA': 'Nbase_A', 'NbC': 'Nbase_C'}.get(k, k)])
# R1 load derivation + per-combination peak rafter moment (Combos sheet)
for k, pk in (("l_pm", L["pm"]), ("l_trib", L["trib"]), ("l_udl_m", L["wm"]), ("l_udl_r", L["wr"]), ("g_sa", g["sa"]), ("g_sb", g["sb"]), ("g_Lm", g["Lm"]), ("g_BCdwg", g["BCdwg"]), ("g_mod_ovh", g["mod_ovh"])):
    add(k, val(k), pk)
add("l_udl_chk", val("l_udl_chk"), 0.0)
cb = wb["Combos"]; cr = bk.reg["Combos"]
for j, (nm, _, kind) in enumerate(E.COMBOS):
    c_ = o["cmb"][nm]
    for k in ("Mstar", "sstar"): add(f"Combos {nm[:2]} {k}", cb.cell(cr[k], 4 + j).value, c_[k])
# member / connection / ballast / anchorage / deflection
add("m_Td", val("m_Td"), mc["Td"]); add("m_Tdn", val("m_Tdn"), mc["Tdn"]); add("m_beta", val("m_beta"), mc["beta"])
for s_, k in (("v", "vert"), ("i", "inc"), ("b", "base")): add(f"m_Pd_{s_}", val(f"m_Pd_{s_}"), mc[f"Pd_{k}"])
add("m_Mdg", val("m_Mdg"), mc["Mdg"]); add("m_Mdn", val("m_Mdn"), mc["Mdn"]); add("m_Vd", val("m_Vd"), mc["Vd"])
ch = o["chk_member"]
add("mk_vT", val("mk_vT_ur"), ch["Vertical member  N (T)"][0]/ch["Vertical member  N (T)"][1]); add("mk_vC", val("mk_vC_ur"), ch["Vertical member  N (C)"][0]/ch["Vertical member  N (C)"][1])
add("mk_iT", val("mk_iT_ur"), ch["Inclined  N+M (T)  cl 9.3.1"][0]); add("mk_iC", val("mk_iC_ur"), ch["Inclined  N+M (C)  cl 9.3.1"][0])
add("mk_iV", val("mk_iV_ur"), ch["Inclined  shear cl 8.4"][0]/ch["Inclined  shear cl 8.4"][1])
add("mk_bS", val("mk_bS_ur"), max(ch["Base  N+M at J-bolt section (T)"][0], ch["Base  N+M at J-bolt section (C)"][0]))
add("mk_bM", val("mk_bM_ur"), ch["Base  N+M at mid-span"][0]); add("mk_bV", val("mk_bV_ur"), ch["Base  shear cl 8.4"][0]/ch["Base  shear cl 8.4"][1])
for j, k in (("A", "M10 A base-vertical"), ("B", "M10 B vertical-inclined"), ("C", "M12 C base-inclined"), ("M", "M8 module (SS A2-70)")):
    add(f"ck_{j}", val(f"ck_{j}_ur"), o["chk_bolt"][k]["ur"]); add(f"c{j}_Vdpb", val(f"c{j}_Vdpb"), o["chk_bolt"][k]["Vdpb"]); add(f"c{j}_Tdb", val(f"c{j}_Tdb"), o["chk_bolt"][k]["Tdb"])
for k, nm in (("b_Wb_up", "Wb_uplift"), ("b_Wb_sl", "Wb_slide"), ("b_Wb_ot", "Wb_overt"), ("b_Wb_bk", "Wb_block"), ("b_Wb_req", "Wb_req"), ("b_L_req", "L_req"), ("b_L_rnd", "L_req_round"), ("b_Wb1", "Wb1")):
    add(k, val(k), bal[nm])
for k, nm in (("bk_up", "uplift"), ("bk_sl", "slide"), ("bk_ot", "overt"), ("bk_bl", "block_up"), ("bk_sd", "slide_dn"), ("bk_od", "overt_dn"), ("bk_se", "slide_eq"), ("bk_oe", "overt_eq")):
    d, c = bal[nm]; add(k + "_ur", val(k + "_ur"), d/c)
add("b_roof", val("b_roof"), bal["roof_load_kN_m2"])
add("a_T", val("a_T"), o["anchor"]["M16"]["T"]); add("a_V", val("a_V"), o["anchor"]["M16"]["V"])
for nm in ("M16", "M10"):
    a = o["anchor"][nm]
    add(f"ak_{nm}_T", val(f"ak_{nm}_T_ur"), a["T"]/a["Tdb"]); add(f"ak_{nm}_C", val(f"ak_{nm}_C_ur"), a["comb"]); add(f"ak_{nm}_B", val(f"ak_{nm}_B_ur"), a["T"]/a["Tanch"])
    add(f"a_{nm}_Lproj", val(f"a_{nm}_Lproj"), a["Lproj"]); add(f"a_{nm}_Ltot", val(f"a_{nm}_Ltot"), a["Ltot"]); add(f"a_{nm}_Lsreq", val(f"a_{nm}_Lsreq"), a["Ls_req"])
dl, lim = o["defl"], o["defl_lim"]
add("dk_inc", val("dk_inc_ur"), max(dl[n]["inc"] for n in dl)/lim["inc"]); add("dk_A", val("dk_A_ur"), max(dl[n]["dA"] for n in dl)/lim["tip_A"]); add("dk_C", val("dk_C_ur"), max(dl[n]["dC"] for n in dl)/lim["tip_C"])
worst, bad = 0.0, []
for lbl, x, y in cmp:
    if x is None: bad.append((lbl, x, y)); continue
    err = abs(x - y)/max(1e-9, abs(y)) if abs(y) > 1e-9 else abs(x - y)
    worst = max(worst, err)
    if err > 1e-6: bad.append((lbl, x, y))
print("compared %d cells; worst relative difference %.2e" % (len(cmp), worst))
print("MISMATCHES:", bad if bad else "NONE")
# equilibrium self-checks in the workbook itself
for j, lc in enumerate(BX.LCS):
    print(lc, "eqV=%.2e  eqN=%.2e  eqM=%.2e" % (an.cell(ar["eqV"], 4 + j).value, an.cell(ar["eqN"], 4 + j).value, an.cell(ar["eqM"], 4 + j).value))
print("Summary: overall =", [c.value for row in wb["Summary"].iter_rows() for c in row if c.value in ("PASS", "REVIEW") and c.column == 7][-1:],
      "| fails:", sum(1 for row in wb["Summary"].iter_rows() for c in row if c.value == "FAIL"))
