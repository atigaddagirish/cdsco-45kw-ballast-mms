"""Change inputs in the workbook -> LibreOffice recalculation -> compare with the engine run on the same inputs."""
import subprocess, os, shutil, sys
from openpyxl import load_workbook
import engine as E
import build_excel as BX

T = "/tmp/claude-0/-home-user-Cluade/cb1ede1f-45e1-5aee-bb0b-19ef0d08b619/scratchpad/linked"
shutil.rmtree(T, ignore_errors=True); os.makedirs(T)
XL = "../excel/CDSCO_45kW_Ballast_MMS_Design_Calc.xlsx"
bk = BX.build(os.path.join(T, "probe.xlsx"))
scen = {"k1_risk": 1.0, "terrain": 3, "mu_f": 0.5, "block_L": 700.0, "mod_kg": 27.5, "Vb": 39.0, "z_bldg": 16.0, "wind_mode": 0, "mod_W": 1100.0, "frame_sp": 1350.0, "tilt_probe": None}
scen.pop("tilt_probe")
wb = load_workbook(XL)
for k, v in scen.items():
    sh, ref = bk.names[k].split("!"); wb[sh][ref.replace("$", "")].value = v
mod = os.path.join(T, "modified.xlsx"); wb.save(mod)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(T, "out"), mod], check=True, capture_output=True, timeout=180)
r = load_workbook(os.path.join(T, "out", "modified.xlsx"), data_only=True)
val = lambda n: r[bk.names[n].split("!")[0]][bk.names[n].split("!")[1].replace("$", "")].value
p = E.P0(); p.update(scen); o = E.design(p); b = o["bal"]; w = o["w"]
checks = [("w_k2", w["k2"]), ("w_pd", w["pd"]), ("w_Nup", w["N_up"]), ("b_Wb1", b["Wb1"]), ("b_L_req", b["L_req"]), ("b_Wb_sl", b["Wb_slide"]),
          ("bk_up_ur", b["uplift"][0]/b["uplift"][1]), ("bk_sl_ur", b["slide"][0]/b["slide"][1]), ("bk_ot_ur", b["overt"][0]/b["overt"][1]), ("bk_bl_ur", b["block_up"][0]/b["block_up"][1]),
          ("m_iM", o["env"]["inc_M"]), ("mk_iT_ur", o["chk_member"]["Inclined  N+M (T)  cl 9.3.1"][0]), ("l_udl_m", o["L"]["wm"]), ("a_T", o["anchor"]["M10"]["T"]), ("a_M10_Ltot", o["anchor"]["M10"]["Ltot"]), ("l_Wm", o["L"]["Wm"]),
          ("mk_bS_ur", max(o["chk_member"]["Base  N+M at J-bolt section (T)"][0], o["chk_member"]["Base  N+M at J-bolt section (C)"][0]))]
worst = 0
for n, y in checks:
    x = val(n); e = abs(x - y)/max(1e-9, abs(y)); worst = max(worst, e); print("%-12s excel %.6g  python %.6g  rel.diff %.1e" % (n, x, y, e))
print("inputs changed:", scen); print("WORST %.1e -> %s" % (worst, "LIVE LINKAGE CONFIRMED" if worst < 1e-6 else "MISMATCH"))
